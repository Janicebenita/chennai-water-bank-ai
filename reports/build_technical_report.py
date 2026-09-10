"""Build the Chennai Water Bank technical report DOCX from the audited Markdown.

Design preset: compact_reference_guide.
Header template: editorial_cover.
Named overrides:
- editorial cover uses larger display typography;
- evidence tables with six or more columns use Letter landscape with 0.5-inch
  margins to retain readable text.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Chennai_Water_Bank_Technical_Development_Report.md"
OUTPUT = ROOT / "Chennai_Water_Bank_Technical_Development_Report.docx"

NAVY = "17365D"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
TEAL = "138D86"
TEAL_LIGHT = "DFF3F1"
AMBER = "F2C36B"
AMBER_LIGHT = "FFF5DF"
TEXT = "263238"
MUTED = "5D6D73"
LIGHT = "E8EEF5"
VERY_LIGHT = "F6F8FA"
WHITE = "FFFFFF"
CODE_BG = "F1F4F6"
GRID = "AEBCC5"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, total_width_twips: int, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_width_twips))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            set_cell_margins(cell)


def add_page_field(paragraph) -> None:
    paragraph.add_run("PAGE ").font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])


def set_paragraph_border(paragraph, color: str, side: str = "left", size: str = "18") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "8")
    border.set(qn("w:color"), color)
    p_bdr.append(border)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_inline_runs(paragraph, text: str, *, color: str | None = None, size: float | None = None) -> None:
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(size or 9)
            run.font.color.rgb = RGBColor.from_string(BLUE_DARK)
        if color and not token.startswith("`"):
            run.font.color.rgb = RGBColor.from_string(color)
        if size and not token.startswith("`"):
            run.font.size = Pt(size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(BLUE)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.page_break_before = True

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(7)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(5)
    h3.paragraph_format.keep_with_next = True

    for list_style_name in ("List Bullet", "List Number"):
        style = styles[list_style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        style.paragraph_format.line_spacing = 1.25

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor.from_string("26343A")
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.12)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(6)
    code_style.paragraph_format.line_spacing = 1.0

    if "Evidence Callout" not in styles:
        callout = styles.add_style("Evidence Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Evidence Callout"]
    callout.font.name = "Calibri"
    callout.font.size = Pt(10)
    callout.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.08)
    callout.paragraph_format.space_before = Pt(4)
    callout.paragraph_format.space_after = Pt(7)
    callout.paragraph_format.line_spacing = 1.15


def configure_section(section, *, landscape: bool = False, first: bool = False) -> None:
    section.top_margin = Inches(1.0 if not landscape else 0.5)
    section.bottom_margin = Inches(1.0 if not landscape else 0.5)
    section.left_margin = Inches(1.0 if not landscape else 0.5)
    section.right_margin = Inches(1.0 if not landscape else 0.5)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    section.different_first_page_header_footer = first


def configure_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = True
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    left.width = Inches(3.3)
    right.width = Inches(3.2)
    for cell in (left, right):
        set_cell_margins(cell, top=0, bottom=25, start=0, end=0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("CHENNAI WATER BANK")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("TECHNICAL DEVELOPMENT REPORT")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    line = header.add_paragraph()
    line.paragraph_format.space_after = Pt(0)
    line.paragraph_format.space_before = Pt(0)
    set_paragraph_border(line, TEAL, side="bottom", size="8")

    footer = section.footer
    footer.is_linked_to_previous = True
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("SIMULATION-DRIVEN PROTOTYPE  •  EVIDENCE-BASED REPORT  •  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_field(p)


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    configure_section(section, first=True)

    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    band.autofit = False
    set_table_geometry(band, 9360, [9360])
    cell = band.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=180, bottom=180, start=220, end=220)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("WATER • DATA • DECISION • CONTROL")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(WHITE)

    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("CHENNAI WATER BANK")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("Evidence-Based Technical Development Write-up\nand Software Engineering Report")
    run.font.name = "Calibri"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("ACADEMIC • VIVA • ARCHITECTURE • MAINTENANCE")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    meta = doc.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    set_table_geometry(meta, 7800, [2100, 5700])
    metadata = [
        ("AUDIT DATE", "17 August 2026"),
        ("CODEBASE", r"C:\Users\User\Documents\ChatGPT\Chennai Water Bank"),
        ("REVISION", "Initial uncommitted worktree — no commit hash available"),
        ("AUTHOR / INSTITUTION", "Not specified in the repository"),
        ("DELIVERABLE", "Complete repository-audited engineering report"),
    ]
    for idx, (label, value) in enumerate(metadata):
        for cell in meta.rows[idx].cells:
            set_cell_margins(cell, top=90, bottom=90, start=120, end=120)
            set_cell_shading(cell, VERY_LIGHT if idx % 2 == 0 else WHITE)
        p = meta.cell(idx, 0).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(TEAL)
        p = meta.cell(idx, 1).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(value)
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(TEXT)

    doc.add_paragraph()
    warning = doc.add_paragraph(style="Evidence Callout")
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.paragraph_format.left_indent = Inches(0.55)
    warning.paragraph_format.right_indent = Inches(0.55)
    set_paragraph_shading(warning, AMBER_LIGHT)
    set_paragraph_border(warning, AMBER)
    run = warning.add_run(
        "SIMULATION-DRIVEN PROTOTYPE — NO PHYSICAL IOT SENSOR OR ACTUATOR CONNECTED"
    )
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Prepared from the current source code, configuration, tests, data files, and available runtime evidence")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    doc.add_page_break()


def paragraph_with_text(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    add_inline_runs(paragraph, text)
    if text.startswith("**VERIFIED FROM") or text.startswith("**INFERRED FROM") or text.startswith("**PROPOSED FUTURE"):
        paragraph.style = doc.styles["Evidence Callout"]
        set_paragraph_shading(paragraph, TEAL_LIGHT)
        set_paragraph_border(paragraph, TEAL)
    return paragraph


def table_column_widths(headers: list[str], rows: list[list[str]], total_twips: int) -> list[int]:
    weights = []
    for idx, header in enumerate(headers):
        values = [header] + [row[idx] if idx < len(row) else "" for row in rows]
        longest = max(len(value) for value in values)
        weights.append(max(7, min(32, longest)))
    weight_sum = sum(weights)
    widths = [max(700, round(total_twips * weight / weight_sum)) for weight in weights]
    delta = total_twips - sum(widths)
    widths[-1] += delta
    return widths


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> bool:
    wide = len(headers) >= 6
    if wide:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, landscape=True)
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
        total_twips = 14400
        font_size = 8.0
    else:
        total_twips = 9360
        font_size = 8.5 if len(headers) >= 5 else 9.0

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = table_column_widths(headers, rows, total_twips)

    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, header in enumerate(headers):
        cell = header_row.cells[idx]
        set_cell_shading(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline_runs(p, header, color=NAVY, size=font_size)
        for run in p.runs:
            run.bold = True

    for row_idx, values in enumerate(rows):
        row = table.add_row()
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            if row_idx % 2 == 1:
                set_cell_shading(cell, VERY_LIGHT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline_runs(p, value, color=TEXT, size=font_size)

    set_table_geometry(table, total_twips, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    if wide:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section, landscape=False)
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
    return wide


def parse_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]], int]:
    def split(line: str) -> list[str]:
        return [part.strip() for part in line.strip().strip("|").split("|")]

    headers = split(lines[start])
    rows: list[list[str]] = []
    idx = start + 2
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        values = split(lines[idx])
        if len(values) < len(headers):
            values += [""] * (len(headers) - len(values))
        rows.append(values[: len(headers)])
        idx += 1
    return headers, rows, idx


def render_markdown_body(doc: Document, text: str) -> None:
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## Evidence classification used"))
    lines = lines[start:]
    idx = 0
    in_code = False
    code_lines: list[str] = []

    while idx < len(lines):
        raw = lines[idx]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                paragraph = doc.add_paragraph(style="Code Block")
                set_paragraph_shading(paragraph, CODE_BG)
                set_paragraph_border(paragraph, GRID, size="8")
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                in_code = False
            idx += 1
            continue
        if in_code:
            code_lines.append(raw)
            idx += 1
            continue

        if not stripped:
            idx += 1
            continue
        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            set_paragraph_border(p, LIGHT, side="bottom", size="6")
            idx += 1
            continue

        if stripped.startswith("|") and idx + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[idx + 1]):
            headers, rows, idx = parse_table(lines, idx)
            add_table(doc, headers, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            paragraph = doc.add_heading(title, level=level)
            if level == 1 and title.startswith(("32.", "33.", "34.")):
                paragraph.paragraph_format.page_break_before = True
            idx += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet:
            paragraph_with_text(doc, bullet.group(1), style="List Bullet")
            idx += 1
            continue
        if numbered:
            paragraph_with_text(doc, numbered.group(1), style="List Number")
            idx += 1
            continue
        if stripped.startswith(">"):
            p = paragraph_with_text(doc, stripped.lstrip("> "), style="Evidence Callout")
            p.runs[0].italic = True
            set_paragraph_shading(p, AMBER_LIGHT)
            set_paragraph_border(p, AMBER)
            idx += 1
            continue

        paragraph_with_text(doc, stripped)
        idx += 1


def set_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Chennai Water Bank — Technical Development Report"
    props.subject = "Evidence-based software engineering report for the Chennai Water Bank prototype"
    props.author = "Prepared from audited repository evidence"
    props.keywords = "Chennai Water Bank, Streamlit, simulation, hydrology, decision engine, IoT, PLC"
    props.comments = (
        "Generated with compact_reference_guide design preset and editorial_cover header template."
    )


def main() -> None:
    source_text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_styles(doc)
    set_document_properties(doc)
    add_cover(doc)
    configure_header_footer(doc.sections[0])
    render_markdown_body(doc, source_text)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
